"""
FP&A ETL Pipeline
=================
Extracts the SaaS financial model (Excel) into a normalized MySQL schema
with Budget vs Actual scenarios, ready for Python-side forecasting / PVM
variance analysis and Power BI reporting.

Source workbook : SaaS_Financial_Model_by_Shivam_Kukreja.xlsx
Target database : fpa_project (MySQL)
Tables loaded    : income_statement, balance_sheet, cash_flow,
                   headcount_assumptions, monthly_saas_metrics

Design notes (for interview / README):
  - "Budget" rows are the model's own planned output (ground truth).
  - "Actual" rows are synthesized with realistic, seeded random variance
    on top of Budget, since the model itself has no live actuals.
    This is a standard technique for demonstrating a BvA / flux-analysis
    workflow in a portfolio project and is disclosed as such, not hidden.
  - Derived fields (year, month name, month index) are intentionally NOT
    stored — they're computed on demand from month_date in SQL/Python to
    avoid redundant, driftable data (basic normalization discipline).
"""

from __future__ import annotations

import getpass
import logging
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from sqlalchemy import create_engine
from sqlalchemy.engine import Engine

logging.basicConfig(level=logging.INFO, format="%(levelname)s | %(message)s")
log = logging.getLogger(__name__)

# --------------------------------------------------------------------------
# Configuration
# --------------------------------------------------------------------------

SOURCE_FILE = Path("SaaS_Financial_Model_by_Shivam_Kukreja.xlsx")
PNL_SHEET = "P&L, BS & CF"
METRICS_SHEET = "Monthly Saas Metrics"

MYSQL_HOST = "localhost"
MYSQL_PORT = 3306
MYSQL_DB = "fpa_project"
MYSQL_USER = "root"

RANDOM_SEED = 42          # reproducible "actuals" every run
VARIANCE_RANGE = (0.05, 0.15)   # +/-5% to 15% swing vs Budget

# Row numbers below were read directly off the workbook (P&L, BS & CF tab).
# Mapping is explicit rather than "guessed by position" on purpose: if the
# sheet layout ever changes, this is the one place to fix it.

INCOME_STATEMENT_ROWS = {
    "active_customers": 7,
    "net_bookings": 9,
    "net_revenue_mrr": 11,
    "deferred_rev": 13,
    "annual_run_rate": 15,
    "ltv_cac_ratio": 16,
    "cogs_aws": 20,
    "cogs_customer_support": 22,
    "cogs_benefits_taxes": 23,
    "total_cogs": 25,
    "gross_margin": 27,
    "gm_pct": 28,
    "personnel_sales": 33,
    "personnel_marketing": 34,
    "personnel_creative": 35,
    "personnel_technology": 36,
    "personnel_productmgmt": 37,
    "personnel_execadmin": 38,
    "personnel_bonuses": 39,
    "personnel_benefits_taxes": 40,
    "total_personnel_exp": 41,
    "ad_spend": 44,
    "other_marketing_opex": 45,
    "technology_opex": 46,
    "da_opex": 47,
    "misc_fixed_exp": 48,
    "variable_exp_pct_bookings": 49,
    "total_other_opex": 50,
    "total_opex": 52,
    "ebit": 54,          # "OPERATING PROFIT" row - the dollar EBIT figure
    "ebit_margin_pct": 55,   # "EBIT" row is actually EBIT-as-%-of-revenue in this workbook
    "tax_loss_asset": 57,
    "taxes": 58,
    "net_income": 60,
}

BALANCE_SHEET_ROWS = {
    "cash": 65,
    "fixed_assets": 67,
    "depreciation": 68,
    "net_fixed_assets": 69,
    "total_assets": 71,
    "deferred_revenue_liability": 74,
    "total_liabilities": 76,
    "new_financing": 79,
    "retained_earnings": 80,
    "total_equity": 81,
    "balance_check": 83,
}

CASH_FLOW_ROWS = {
    "net_income": 88,
    "net_deferred_rev": 89,
    "depreciation": 90,
    "operating_cash_flow": 91,
    "capex": 94,
    "new_equity": 97,
    "net_cash_flow": 99,
}

# department label (as it appears in col B) -> row number
HEADCOUNT_ROWS = {
    "Customer Support": 109,
    "Sales": 115,
    "Marketing": 116,
    "Creative": 117,
    "Technology": 118,
    "Product Mgmt": 119,
    "Exec / Admin": 120,
}

METRICS_COLUMN_RENAME = {
    "Customers_Begin": "customers_begin",
    "Customers_New": "customers_new",
    "Customers_Churned": "customers_churned",
    "Customers_End": "customers_end",
    "MRR_Begin": "mrr_begin",
    "Expansion_MRR": "expansion_mrr",
    "Churned_MRR": "churned_mrr",
    "MRR_End": "mrr_end",
    "NDR %": "ndr_pct",
    "COGS": "cogs",
    "Sales_Mkt_Spend": "sales_mkt_spend",
}

FIRST_MONTH_COL = 3   # column C = month index 1
LAST_MONTH_COL = 50   # column AX = month index 48


# --------------------------------------------------------------------------
# Extraction
# --------------------------------------------------------------------------

def load_workbook(path: Path):
    log.info("Opening workbook: %s", path)
    return openpyxl.load_workbook(path, data_only=True)


def get_month_dates(wb) -> list[pd.Timestamp]:
    """Single source of truth for the calendar: pulled from the metrics tab,
    which has real dates, and reused to date-index the P&L tab (which only
    has a month index, 1..48, across its columns)."""
    ws = wb[METRICS_SHEET]
    dates = [
        pd.Timestamp(ws.cell(row=r, column=2).value)
        for r in range(2, ws.max_row + 1)
        if ws.cell(row=r, column=2).value is not None
    ]
    log.info("Found %d months: %s -> %s", len(dates), dates[0].date(), dates[-1].date())
    return dates


def extract_row_block(ws, row_map: dict[str, int], month_dates: list[pd.Timestamp]) -> pd.DataFrame:
    """Pull a set of labeled rows across the month columns into a tidy
    (month_date x field) DataFrame."""
    data = {}
    for field, row in row_map.items():
        values = [
            ws.cell(row=row, column=c).value
            for c in range(FIRST_MONTH_COL, LAST_MONTH_COL + 1)
        ]
        data[field] = values[: len(month_dates)]
    df = pd.DataFrame(data)
    df.insert(0, "month_date", month_dates)
    return df


def extract_headcount(ws, month_dates: list[pd.Timestamp]) -> pd.DataFrame:
    """Long-format: one row per (month, department)."""
    frames = []
    for dept, row in HEADCOUNT_ROWS.items():
        values = [
            ws.cell(row=row, column=c).value
            for c in range(FIRST_MONTH_COL, LAST_MONTH_COL + 1)
        ]
        values = values[: len(month_dates)]
        frames.append(pd.DataFrame({
            "month_date": month_dates,
            "department": dept,
            "headcount": values,
        }))
    return pd.concat(frames, ignore_index=True)


def extract_saas_metrics(wb) -> pd.DataFrame:
    ws = wb[METRICS_SHEET]
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    rows = []
    for r in range(2, ws.max_row + 1):
        row_vals = [ws.cell(row=r, column=c).value for c in range(1, len(headers) + 1)]
        if row_vals[1] is None:   # Date column empty -> end of data
            break
        rows.append(row_vals)
    df = pd.DataFrame(rows, columns=headers)
    df = df.rename(columns=METRICS_COLUMN_RENAME)
    df["month_date"] = pd.to_datetime(df["Date"])
    keep_cols = ["month_date"] + list(METRICS_COLUMN_RENAME.values())
    return df[keep_cols]


# --------------------------------------------------------------------------
# Budget -> synthetic Actuals
# --------------------------------------------------------------------------

def add_actuals(df: pd.DataFrame, key_cols: list[str], seed: int = RANDOM_SEED) -> pd.DataFrame:
    """Given a Budget-only DataFrame, append a same-shaped 'Actual' scenario
    with reproducible randomized variance on numeric columns. Non-numeric /
    key columns are copied through unchanged."""
    rng = np.random.default_rng(seed)
    budget = df.copy()
    budget.insert(1, "scenario", "Budget")

    actual = df.copy()
    numeric_cols = [c for c in actual.columns if c not in key_cols and pd.api.types.is_numeric_dtype(actual[c])]
    for col in numeric_cols:
        pct = rng.uniform(-VARIANCE_RANGE[1], VARIANCE_RANGE[1], size=len(actual))
        # keep sign realistic: skew slightly toward the +/-5-15% band rather than near-zero
        pct = np.where(np.abs(pct) < VARIANCE_RANGE[0], np.sign(pct) * VARIANCE_RANGE[0], pct)
        actual[col] = (actual[col].astype(float) * (1 + pct)).round(2)
    actual.insert(1, "scenario", "Actual")

    return pd.concat([budget, actual], ignore_index=True)


def add_actuals_headcount(df: pd.DataFrame, seed: int = RANDOM_SEED) -> pd.DataFrame:
    """Headcount varies by whole people, not percentages - separate, simpler
    logic from the financial-line variance above."""
    rng = np.random.default_rng(seed + 1)
    budget = df.copy()
    budget.insert(2, "scenario", "Budget")

    actual = df.copy()
    actual["headcount"] = actual["headcount"].fillna(0)
    delta = rng.integers(-1, 2, size=len(actual))  # -1, 0, or +1 headcount vs plan
    actual["headcount"] = (actual["headcount"].astype(int) + delta).clip(lower=0)
    actual.insert(2, "scenario", "Actual")

    budget["headcount"] = budget["headcount"].fillna(0).astype(int)

    return pd.concat([budget, actual], ignore_index=True)


# --------------------------------------------------------------------------
# Load
# --------------------------------------------------------------------------

def get_engine(password: str) -> Engine:
    url = f"mysql+pymysql://{MYSQL_USER}:{password}@{MYSQL_HOST}:{MYSQL_PORT}/{MYSQL_DB}"
    return create_engine(url)


def load_table(engine: Engine, df: pd.DataFrame, table_name: str) -> None:
    log.info("Loading %d rows into %s", len(df), table_name)
    df.to_sql(table_name, engine, if_exists="append", index=False, method="multi", chunksize=500)


# --------------------------------------------------------------------------
# Orchestration
# --------------------------------------------------------------------------

def main() -> None:
    wb = load_workbook(SOURCE_FILE)
    pnl_ws = wb[PNL_SHEET]
    month_dates = get_month_dates(wb)

    income_stmt = extract_row_block(pnl_ws, INCOME_STATEMENT_ROWS, month_dates)
    balance_sheet = extract_row_block(pnl_ws, BALANCE_SHEET_ROWS, month_dates)
    cash_flow = extract_row_block(pnl_ws, CASH_FLOW_ROWS, month_dates)
    headcount = extract_headcount(pnl_ws, month_dates)
    saas_metrics = extract_saas_metrics(wb)

    income_stmt = add_actuals(income_stmt, key_cols=["month_date"])
    balance_sheet = add_actuals(balance_sheet, key_cols=["month_date"])
    cash_flow = add_actuals(cash_flow, key_cols=["month_date"])
    headcount = add_actuals_headcount(headcount)
    saas_metrics = add_actuals(saas_metrics, key_cols=["month_date"])

    password = getpass.getpass("MySQL root password: ")
    engine = get_engine(password)

    load_table(engine, income_stmt, "income_statement")
    load_table(engine, balance_sheet, "balance_sheet")
    load_table(engine, cash_flow, "cash_flow")
    load_table(engine, headcount, "headcount_assumptions")
    load_table(engine, saas_metrics, "monthly_saas_metrics")

    log.info("Done. All 5 tables loaded with Budget + Actual scenarios.")


if __name__ == "__main__":
    main()
